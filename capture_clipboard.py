"""Watch the Windows clipboard and append everything you copy to a Word document.

Run it, then go about your browsing. Every time you press Ctrl+A / Ctrl+C on a
page (or copy a single image), this appends the content -- text and pictures --
as a new section of a .docx file. Press Ctrl+C in the console window to stop.

Usage:
    python capture_clipboard.py
    python capture_clipboard.py --out "C:\\Users\\me\\Desktop\\research.docx"

Requires: pywin32, Pillow, python-docx  (see requirements.txt)
"""

from __future__ import annotations

import argparse
import base64
import binascii
import datetime as _dt
import hashlib
import io
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser

try:
    import win32clipboard
except ImportError:  # pragma: no cover - dependency hint
    sys.exit("pywin32 is missing. Run:  pip install -r requirements.txt")

try:
    from PIL import Image, ImageGrab
except ImportError:  # pragma: no cover
    sys.exit("Pillow is missing. Run:  pip install -r requirements.txt")

try:
    from docx import Document
    from docx.shared import Inches, Pt
except ImportError:  # pragma: no cover
    sys.exit("python-docx is missing. Run:  pip install -r requirements.txt")


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)
MAX_IMAGE_BYTES = 25 * 1024 * 1024
PAGE_WIDTH_INCHES = 6.0
EMBEDDABLE = {"PNG", "JPEG", "GIF", "BMP", "TIFF"}


# --------------------------------------------------------------------------
# Clipboard reading
# --------------------------------------------------------------------------

def _open_clipboard(attempts: int = 12, delay: float = 0.08) -> bool:
    """The clipboard is a single shared resource; another app may hold it."""
    for _ in range(attempts):
        try:
            win32clipboard.OpenClipboard()
            return True
        except Exception:
            time.sleep(delay)
    return False


def read_clipboard() -> dict:
    """Return {'text': str|None, 'html': str|None, 'source_url': str|None, 'bitmap': bytes|None}."""
    out = {"text": None, "html": None, "source_url": None, "bitmap": None}
    if not _open_clipboard():
        return out
    try:
        html_fmt = win32clipboard.RegisterClipboardFormat("HTML Format")
        if win32clipboard.IsClipboardFormatAvailable(html_fmt):
            try:
                raw = win32clipboard.GetClipboardData(html_fmt)
                if isinstance(raw, bytes):
                    raw = raw.decode("utf-8", "replace")
                out["html"], out["source_url"] = split_cf_html(raw)
            except Exception:
                pass
        if win32clipboard.IsClipboardFormatAvailable(win32clipboard.CF_UNICODETEXT):
            try:
                out["text"] = win32clipboard.GetClipboardData(win32clipboard.CF_UNICODETEXT)
            except Exception:
                pass
        has_bitmap = win32clipboard.IsClipboardFormatAvailable(win32clipboard.CF_DIB)
    finally:
        win32clipboard.CloseClipboard()

    if has_bitmap:
        # ImageGrab does the DIB -> PIL conversion for us, and it needs the
        # clipboard closed again before it opens it itself.
        try:
            grabbed = ImageGrab.grabclipboard()
        except Exception:
            grabbed = None
        if isinstance(grabbed, Image.Image):
            buf = io.BytesIO()
            grabbed.convert("RGB").save(buf, format="PNG")
            out["bitmap"] = buf.getvalue()
    return out


def split_cf_html(raw: str) -> tuple[str, str | None]:
    """Strip the CF_HTML header off clipboard HTML and pull out the SourceURL."""
    source = None
    m = re.search(r"^SourceURL:(.*)$", raw, re.MULTILINE)
    if m:
        source = m.group(1).strip() or None

    start = raw.find("<!--StartFragment-->")
    end = raw.find("<!--EndFragment-->")
    if start != -1 and end > start:
        return raw[start + len("<!--StartFragment-->"):end], source

    # No fragment markers: drop the header lines and keep the rest.
    body = re.sub(r"^(?:Version|StartHTML|EndHTML|StartFragment|EndFragment|SourceURL)"
                  r":.*$\n?", "", raw, flags=re.MULTILINE)
    return body, source


# --------------------------------------------------------------------------
# HTML -> flat list of document operations
# --------------------------------------------------------------------------

class PageParser(HTMLParser):
    """Flatten HTML into ordered ops: ('h', level, text) ('p', text) ('li', text, depth) ('img', url, alt)."""

    SKIP = {"script", "style", "head", "noscript", "svg", "iframe", "object", "template", "select"}
    BLOCK = {
        "p", "div", "br", "hr", "blockquote", "pre", "section", "article", "header",
        "footer", "main", "aside", "nav", "ul", "ol", "dl", "dd", "dt", "figure",
        "figcaption", "form", "table", "caption", "address", "details", "summary",
    }
    HEADINGS = {f"h{n}": n for n in range(1, 7)}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.ops: list[tuple] = []
        self._buf: list[str] = []
        self._skip = 0
        self._heading: int | None = None
        self._list_depth = 0
        self._in_li = False
        self._in_pre = False
        self._in_row = False
        self._cells: list[str] = []

    # -- helpers ----------------------------------------------------------
    def _text(self) -> str:
        text = "".join(self._buf)
        self._buf.clear()
        if self._in_pre:
            return text.strip("\n").rstrip()
        return re.sub(r"\s+", " ", text).strip()

    def _flush(self) -> None:
        text = self._text()
        if self._in_row:
            if text:
                self._cells.append(text)
            return
        if not text:
            self._in_li = False
            return
        if self._heading is not None:
            self.ops.append(("h", self._heading, text))
        elif self._in_li:
            self.ops.append(("li", text, max(1, self._list_depth)))
        else:
            self.ops.append(("p", text))
        self._in_li = False

    def _end_row(self) -> None:
        self._flush()
        if self._cells:
            self.ops.append(("p", "  |  ".join(self._cells)))
        self._cells = []
        self._in_row = False

    # -- HTMLParser hooks -------------------------------------------------
    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
            return
        if self._skip:
            return
        attrd = {k.lower(): (v or "") for k, v in attrs}

        if tag == "img":
            self._flush()
            src = attrd.get("src") or _first_srcset(attrd.get("srcset", ""))
            if src:
                self.ops.append(("img", src.strip(), attrd.get("alt", "").strip()))
            return
        if tag == "tr":
            self._flush()
            self._in_row = True
            self._cells = []
            return
        if tag in ("td", "th"):
            self._flush()
            return
        if tag in self.HEADINGS:
            self._flush()
            self._heading = self.HEADINGS[tag]
            return
        if tag in ("ul", "ol"):
            self._flush()
            self._list_depth += 1
            return
        if tag == "li":
            self._flush()
            self._in_li = True
            return
        if tag == "pre":
            self._flush()
            self._in_pre = True
            return
        if tag in self.BLOCK:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
            return
        if self._skip:
            return
        if tag == "tr":
            self._end_row()
            return
        if tag in ("td", "th"):
            self._flush()
            return
        if tag in self.HEADINGS:
            self._flush()
            self._heading = None
            return
        if tag in ("ul", "ol"):
            self._flush()
            self._list_depth = max(0, self._list_depth - 1)
            return
        if tag == "pre":
            self._flush()
            self._in_pre = False
            return
        if tag in self.BLOCK or tag == "li":
            self._flush()

    def handle_data(self, data):
        if self._skip:
            return
        self._buf.append(data)

    def close(self):
        super().close()
        if self._in_row:
            self._end_row()
        else:
            self._flush()


def _first_srcset(srcset: str) -> str:
    for candidate in srcset.split(","):
        url = candidate.strip().split(" ")[0]
        if url:
            return url
    return ""


def ops_from_html(html: str, source_url: str | None) -> list[tuple]:
    parser = PageParser()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        pass
    ops = []
    for op in parser.ops:
        if op[0] == "img" and source_url:
            op = ("img", urllib.parse.urljoin(source_url, op[1]), op[2])
        ops.append(op)
    return ops


def ops_from_text(text: str) -> list[tuple]:
    ops = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        for line in block.split("\n"):
            line = line.rstrip()
            if line.strip():
                ops.append(("p", line))
    return ops


# --------------------------------------------------------------------------
# Image fetching
# --------------------------------------------------------------------------

class ImageFetcher:
    def __init__(self, referer: str | None, min_size: int, timeout: float) -> None:
        self.referer = referer
        self.min_size = min_size
        self.timeout = timeout
        self.cache: dict[str, bytes | None] = {}

    def fetch(self, url: str) -> bytes | None:
        if url in self.cache:
            return self.cache[url]
        data = self._fetch_raw(url)
        self.cache[url] = data
        return data

    def _fetch_raw(self, url: str) -> bytes | None:
        if url.startswith("data:"):
            return _decode_data_uri(url)
        if not url.lower().startswith(("http://", "https://")):
            return None
        headers = {"User-Agent": USER_AGENT, "Accept": "image/*,*/*"}
        if self.referer:
            headers["Referer"] = self.referer
        req = urllib.request.Request(url, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                return resp.read(MAX_IMAGE_BYTES + 1)
        except (urllib.error.URLError, urllib.error.HTTPError, OSError, ValueError):
            return None

    def prepare(self, data: bytes) -> tuple[bytes, float, float] | None:
        """Normalise to something Word accepts; return (bytes, width_in, height_in)."""
        if not data or len(data) > MAX_IMAGE_BYTES:
            return None
        try:
            img = Image.open(io.BytesIO(data))
            img.load()
        except Exception:
            return None
        if img.width < self.min_size or img.height < self.min_size:
            return None  # spacer / tracking pixel / tiny icon

        fmt = (img.format or "").upper()
        if fmt not in EMBEDDABLE:
            buf = io.BytesIO()
            img.convert("RGBA" if img.mode in ("RGBA", "LA", "P") else "RGB").save(buf, format="PNG")
            data = buf.getvalue()

        width_in = img.width / 96.0
        height_in = img.height / 96.0
        if width_in > PAGE_WIDTH_INCHES:
            height_in *= PAGE_WIDTH_INCHES / width_in
            width_in = PAGE_WIDTH_INCHES
        return data, width_in, height_in


def _decode_data_uri(url: str) -> bytes | None:
    try:
        header, _, payload = url.partition(",")
        if not payload:
            return None
        if ";base64" in header.lower():
            return base64.b64decode(payload, validate=False)
        return urllib.parse.unquote_to_bytes(payload)
    except (binascii.Error, ValueError):
        return None


# --------------------------------------------------------------------------
# Document building
# --------------------------------------------------------------------------

class CaptureDoc:
    def __init__(self, path: str) -> None:
        self.path = path
        self.doc = Document()
        self.count = 0
        self.images = 0
        self.pending_save = False
        self.doc.add_heading("Clipboard capture", level=0)
        stamp = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        intro = self.doc.add_paragraph(f"Session started {stamp}")
        intro.runs[0].font.size = Pt(9)

    def add_capture(self, ops: list[tuple], source_url: str | None, fetcher: ImageFetcher) -> int:
        self.count += 1
        stamp = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if self.count > 1:
            self.doc.add_page_break()
        self.doc.add_heading(f"Capture {self.count} - {stamp}", level=1)
        if source_url:
            p = self.doc.add_paragraph(source_url)
            p.runs[0].font.size = Pt(9)

        added_images = 0
        for op in ops:
            kind = op[0]
            if kind == "h":
                self.doc.add_heading(op[2], level=min(op[1] + 1, 9))
            elif kind == "p":
                self.doc.add_paragraph(op[1])
            elif kind == "li":
                style = "List Bullet" if op[2] <= 1 else f"List Bullet {min(op[2], 3)}"
                try:
                    self.doc.add_paragraph(op[1], style=style)
                except KeyError:
                    self.doc.add_paragraph(f"- {op[1]}")
            elif kind == "img":
                if self._add_image(op[1], op[2], fetcher):
                    added_images += 1
            elif kind == "rawimg":
                if self._add_raw_image(op[1], fetcher):
                    added_images += 1
        self.images += added_images
        self.save()
        return added_images

    def _add_image(self, url: str, alt: str, fetcher: ImageFetcher) -> bool:
        raw = fetcher.fetch(url)
        prepared = fetcher.prepare(raw) if raw else None
        if not prepared:
            if alt:
                note = self.doc.add_paragraph(f"[image not saved: {alt}]")
                note.runs[0].font.size = Pt(9)
            return False
        return self._place(prepared, alt)

    def _add_raw_image(self, data: bytes, fetcher: ImageFetcher) -> bool:
        prepared = fetcher.prepare(data)
        return self._place(prepared, "") if prepared else False

    def _place(self, prepared: tuple[bytes, float, float], alt: str) -> bool:
        data, width_in, _ = prepared
        try:
            self.doc.add_picture(io.BytesIO(data), width=Inches(width_in))
        except Exception:
            return False
        if alt:
            cap = self.doc.add_paragraph(alt)
            cap.runs[0].font.size = Pt(8)
        return True

    def save(self) -> None:
        """Save after every capture so nothing is lost if the script is killed."""
        try:
            self.doc.save(self.path)
            self.pending_save = False
        except PermissionError:
            self.pending_save = True
            print(f"  ! Could not write {self.path} (is it open in Word?). "
                  f"Everything is still held in memory and will be saved on the next capture.")
        except Exception as exc:
            self.pending_save = True
            print(f"  ! Could not write {self.path}: {exc}")


# --------------------------------------------------------------------------
# Main loop
# --------------------------------------------------------------------------

def default_out_path() -> str:
    stamp = _dt.datetime.now().strftime("%Y-%m-%d_%H%M")
    folder = os.path.join(os.path.expanduser("~"), "Documents")
    if not os.path.isdir(folder):
        folder = os.path.expanduser("~")
    return os.path.join(folder, f"clipboard-capture_{stamp}.docx")


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", "-o", default=None,
                    help="output .docx path (default: ~/Documents/clipboard-capture_<date>.docx)")
    ap.add_argument("--interval", type=float, default=0.4,
                    help="seconds between clipboard checks (default 0.4)")
    ap.add_argument("--no-images", action="store_true",
                    help="capture text only; do not download or embed pictures")
    ap.add_argument("--min-image-size", type=int, default=32,
                    help="skip images smaller than this many pixels on a side (default 32)")
    ap.add_argument("--image-timeout", type=float, default=10.0,
                    help="seconds to wait for each image download (default 10)")
    ap.add_argument("--allow-duplicates", action="store_true",
                    help="record a capture even if it is identical to the previous one")
    ap.add_argument("--min-chars", type=int, default=1,
                    help="ignore text copies shorter than this many characters (default 1)")
    return ap.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    if os.name != "nt":
        print("This script uses the Windows clipboard API and only runs on Windows.")
        return 1

    out_path = os.path.abspath(args.out or default_out_path())
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    doc = CaptureDoc(out_path)

    print("Clipboard capture running.")
    print(f"  Writing to: {out_path}")
    print("  Copy anything (Ctrl+A then Ctrl+C on a page) and it gets appended.")
    print("  Press Ctrl+C here to stop.\n")

    last_seq = win32clipboard.GetClipboardSequenceNumber()
    last_hash = None

    try:
        while True:
            time.sleep(args.interval)
            seq = win32clipboard.GetClipboardSequenceNumber()
            if seq == last_seq:
                continue
            last_seq = seq
            time.sleep(0.12)  # let the copying app finish filling every format

            clip = read_clipboard()
            ops: list[tuple] = []
            source_url = clip["source_url"]

            if clip["html"]:
                ops = ops_from_html(clip["html"], source_url)
            elif clip["text"]:
                ops = ops_from_text(clip["text"])

            if clip["bitmap"] and not args.no_images:
                ops.append(("rawimg", clip["bitmap"]))

            if args.no_images:
                ops = [op for op in ops if op[0] not in ("img", "rawimg")]

            text_len = sum(len(op[1]) for op in ops if op[0] == "p") + \
                sum(len(op[2]) for op in ops if op[0] == "h") + \
                sum(len(op[1]) for op in ops if op[0] == "li")
            has_image = any(op[0] in ("img", "rawimg") for op in ops)
            if not ops or (text_len < args.min_chars and not has_image):
                continue

            fingerprint = hashlib.sha256(
                repr([op if op[0] != "rawimg" else ("rawimg", len(op[1])) for op in ops])
                .encode("utf-8", "replace")
            ).hexdigest()
            if fingerprint == last_hash and not args.allow_duplicates:
                continue
            last_hash = fingerprint

            fetcher = ImageFetcher(
                referer=source_url,
                min_size=args.min_image_size,
                timeout=args.image_timeout,
            )
            n_img = doc.add_capture(ops, source_url, fetcher)
            label = source_url or ("image" if has_image and text_len == 0 else "clipboard text")
            print(f"[{doc.count}] {text_len} chars, {n_img} image(s)  <- {label}")
    except KeyboardInterrupt:
        print("\nStopping.")
    finally:
        if doc.count:
            if doc.pending_save:
                doc.save()
            print(f"Saved {doc.count} capture(s) and {doc.images} image(s) to:\n  {out_path}")
        else:
            print("Nothing was captured, so no file was written.")
            try:
                if os.path.exists(out_path) and doc.count == 0:
                    os.remove(out_path)
            except OSError:
                pass
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
